"""
读取 病历质控-规则需求.xlsx 的「住院最新核对版」页签，
筛选 质控维度=规则质控 的行，根据「判断依据」自动匹配现有计算符，
输出 [判断依据, 计算符, 判断数据需求, 升级建议] 到新的 xlsx。
"""
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

SRC = 'docs/病历质控-规则需求.xlsx'
DST = 'docs/住院规则-计算符匹配.xlsx'

# 现有计算符（来自 DrlCompiler.java）
EXIST_OPS = {
    'eq', 'ne', 'gt', 'lt', 'ge', 'le', 'between', 'IN_SET', 'contains',
    'regex_match', 'regex_not_match', 'regexMatch', 'multiRegexMatch',
    'isBlank', 'isNotBlank',
    'dictMatch', 'whitelistMatch', 'existenceConflict',
    'contradictionCheck', 'similarity',
    'arrayContains', 'arrayLength', 'arrayIntersect',
    'dataCheck', 'timeCheck', 'lengthCheck', 'fieldCompare'
}

def classify(basis: str):
    """返回 (operator, 数据需求, 升级建议)"""
    s = (basis or '').strip()
    if not s:
        return ('', '', '')

    # —— 1. 时间约束类：「X 小时/分钟内完成」「X 时间不能早于 Y」 ——
    if re.search(r'(\d+\s*(小时|时|分钟|h|H)\s*内.*(完成|书写|签|创建))', s) \
       or re.search(r'\d+\s*(天|周|日)\s*内.*完成', s) \
       or re.search(r'\d+\s*(小时|时|分钟|h|H|天|周|日).*(后|内).*(完成|书写)', s):
        return ('timeCheck',
                '需要：被检查文档创建时间字段、参照时间字段（如入院时间、手术结束时间）、阈值（数值+时间单位）',
                '')
    if re.search(r'\d+\s*小时.*(无|缺)', s) or re.search(r'(\d+)\s*小时内(无|缺|入院诊断)', s):
        return ('timeCheck',
                '需要：被检查字段在指定时长内是否存在（如"24小时内无初步诊断"=入院后24h内初步诊断字段为空）',
                '')
    if '时间不能早于' in s or '时间间隔' in s or ('相差' in s and ('小时' in s or '分钟' in s)) \
       or re.search(r'(创建时间|开立时间|签署时间|开嘱时间|开始时间|结束时间).*(-|早于|晚于|>|<).*?(时间|>\s*\d)', s):
        return ('timeCheck',
                '需要：两个时间字段（field、baseTimeField）+ 最小/最大值 + 单位',
                '')
    if re.search(r'每.*天.*查房|每.*超过.*天.*(份|阶段小结)', s):
        return ('timeCheck',
                '需要：起止时间 + 查房记录时间序列 + 间隔阈值（按天）',
                '现有 timeCheck 仅判断单次时间差；建议升级为 periodicTimeCheck（周期性记录补全检查）')

    # —— 2. 非空校验：「X 不为空 / 不能为空 / 需填写 / 需记录 X / 缺 X」 ——
    if re.search(r'不为空|不能为空', s) \
       or re.search(r'需(要|)(填写|有)', s) and not re.search(r'包含|包括|一致|对应|超过|不少于', s) \
       or re.search(r'需记录(?!.*包含|.*包括)', s):
        return ('isNotBlank', '需要：被校验字段', '')
    if re.search(r'^缺[少]?', s) or re.search(r'未书写|未记录$|无.*记录单?$|无.*医嘱$|无.*病程记录$|无.*意见$|缺少', s):
        return ('isNotBlank', '需要：被校验字段（如：术前讨论记录、出院诊断、化疗知情同意书、会诊意见）', '')
    if re.search(r'(需|应当|有).*(签名|签署)', s):
        return ('isNotBlank', '需要：签名字段（如手术者签名、麻醉医师签名、护士签名、主持人审核签名）', '')
    if re.search(r'每项.*有明确', s) and '签名' in s:
        return ('isNotBlank', '需要：医嘱子项签名字段（遍历每条医嘱检查签名是否为空）',
                '现有 isNotBlank 仅检查单个字段；"每项医嘱有明确的签名"需遍历数组，建议升级为 arrayEveryItemNotBlank')
    if re.search(r'每项.*有明确的.*(时间|内容)', s):
        return ('isNotBlank', '需要：医嘱子项时间/内容字段（遍历每条医嘱检查是否为空）',
                '现有 isNotBlank 仅检查单个字段；"每项医嘱有明确的时间"需遍历数组，建议升级为 arrayEveryItemNotBlank')

    # —— 3. 包含/缺项校验：「X 需要包含 Y」「X 内容需包含 Y」「X 包括 Y」 ——
    if re.search(r'(需(要|)(包含|包括|记录并包含)|应当?包(含|括)|内容(需|应当?)(包含|包括))', s):
        return ('contains',
                '需要：父字段（如手术记录、术前小结正文）+ 关键词（如"术前诊断""手术日期""主持人"）',
                '若是结构化子项检查（多个子项任意缺失即扣分），建议升级为 sectionCheck（指定子项数组，逐一 isNotBlank）')

    # —— 4. 字段一致性校验：「X 需与 Y 一致」 ——
    if re.search(r'需?与.*一致', s):
        return ('fieldCompare',
                '需要：两个字段（如 病案首页.过敏史、入院记录.过敏史）+ 比较类型(equals)',
                '现有 fieldCompare 主要做数值/长度比较；建议扩展 compareType=textEquals 与 dictAlias（同义字典归一化后比较）')

    # —— 5. 矛盾性/性别独有：「不能出现女性独有诊断 / 男性独有诊断」 ——
    if '独有诊断' in s or ('不能出现' in s and '诊断' in s):
        return ('existenceConflict',
                '需要：性别字段、诊断信息字段、性别独有诊断字典（如「女性独有诊断字典」「男性独有诊断字典」）',
                '现有 existenceConflict 已支持字典+字段冲突检查，可直接复用')
    if '主诉' in s and '现病史' in s and ('对应' in s or '不一致' in s):
        return ('contradictionCheck',
                '需要：主诉字段、现病史字段、症状字典或矛盾词字典',
                '')
    if '需与' in s and '描述对应' in s:
        return ('contradictionCheck',
                '需要：诊断字段、病史描述字段、症状/疾病关联字典（如"高血压"应对应"头痛""头晕"等症状词）',
                '现有 contradictionCheck 支持两字段矛盾词检查；此规则是正向对应检查，建议扩展 synonymMatch 或 dictCorrelation 计算符')

    # —— 6. 字数限制：「不能超过 X 个汉字」「不少于 X 字」 ——
    if re.search(r'(不(超过|多于|大于)|超过)\s*\d+\s*(个)?(中文)?(汉字|字符|字)', s) \
       or re.search(r'不少于\s*\d+\s*字', s):
        return ('lengthCheck', '需要：被校验文本字段、运算符（< > <= >=）、字数阈值', '')

    # —— 7. 法定结婚年龄等业务规则：年龄+婚姻 ——
    if '法定结婚年龄' in s or ('年龄' in s and '婚姻' in s):
        return ('fieldCompare',
                '需要：性别、年龄、婚姻状况；规则：(性别=男 AND 年龄<22) 或 (性别=女 AND 年龄<20) AND 婚姻∈{已婚,丧偶,离婚}',
                '此规则需多字段联立判断，单 fieldCompare 不够；建议在前端用「与/或」组合多个条件节点（已支持），或升级为 expression 计算符（自由表达式）')

    # —— 8. 既往史/手术史/传染病史关联 ——
    if '既往史' in s and ('手术' in s or '传染病' in s or '外伤' in s):
        return ('contains',
                '需要：既往史字段、关键词列表（如手术、外伤、传染病名称字典）',
                '若需要先判断"患者有手术病史"再校验"既往史是否记录"，则需要两个 contains 条件用 AND 组合')

    # —— 8.5 条件性非空：「有 X 时，缺 Y」「X 有时缺少 Y」 ——
    if re.search(r'(有|存在).*(时|情况下?|时，?)(缺|少|无|缺少)', s) \
       or re.search(r'(有|存在).*(时|，)(缺|少|无)', s):
        return ('isNotBlank',
                '需要：前置条件字段（如会诊医嘱、化疗医嘱、手术医嘱）+ 被校验字段（如会诊单、化疗知情同意书、术前讨论记录），条件不满足时不校验',
                '画布上用两个条件节点组合：条件1=arrayContains(医嘱列表,关键词) AND 条件2=isNotBlank(对应记录)')

    # —— 8.6 书写规范性：「X 书写不规范」 ——
    if '不规范' in s:
        return ('regexMatch',
                '需要：被校验字段 + 正则表达式（如放射剂量格式、放射部位格式）',
                '需业务方提供具体正则规则；现有 regexMatch 可直接使用')

    # —— 9. 输血记录疗效评价 ——
    if '疗效评价' in s and '升高幅度' in s and '输血量' in s:
        return ('expression',
                '需要：输血量、Hb/血小板升高幅度、输血疗效评价文本；规则：升高幅度需匹配输血量预期区间',
                '现有计算符无法处理复杂业务公式；建议升级 expression 计算符（自定义脚本/MVEL 表达式）')

    # —— 10. 医嘱/操作类：医嘱内容包含某类操作 → 病程记录 ——
    if re.search(r'(临时医嘱|医嘱).*(包含|有).*(介入|胸穿|骨穿|输血|化疗|血液制品)', s) \
       and '病程记录' in s:
        return ('contains',
                '需要：临时医嘱列表（数组）、对应病程记录字段、关键词字典',
                '建议组合：先用 arrayContains 判断医嘱中是否含相关操作，再用 contains 校验病程记录子项')

    # —— 10.5 手术患者 + 缺少 X ：条件性非空 ——
    if re.search(r'手术患者.*缺|手术患者.*需有|手术患者.*需', s):
        return ('isNotBlank',
                '需要：前置条件（是否有手术记录/手术医嘱）+ 被校验字段（如术前讨论记录、术者查房记录）',
                '画布上用两个条件节点组合')

    # —— 10.7 全麻患者 + 缺少 X ——
    if re.search(r'全麻.*手术.*需|全麻.*患者.*需', s):
        return ('isNotBlank',
                '需要：前置条件（麻醉方式=全麻）+ 被校验字段（如麻醉术后访视记录单）',
                '画布上用两个条件节点组合')

    # —— 10.8 上级医师查房 48h ——
    if re.search(r'上级医师.*查房.*\d+\s*小时', s):
        return ('timeCheck',
                '需要：入院时间 + 上级医师首次查房记录创建时间 + 阈值(48h)',
                '')

    # —— 10.9 会诊时间差 ——
    if '会诊发出时间' in s and '到达时间' in s and '未记录' in s:
        return ('isNotBlank',
                '需要：会诊发出时间字段、会诊到达时间字段，两者均不得为空',
                '画布上两个 isNotBlank 条件用 AND 组合')

    # —— 10.10 未按要求进行 X 并记录 ——
    if re.search(r'未按要求.*并记录', s):
        return ('isNotBlank',
                '需要：疑难病例讨论记录字段',
                '')

    # —— 10.11 术前常规检验/检查 ——
    if re.search(r'未完成术前常规(检验|检查)', s):
        return ('arrayContains',
                '需要：术前检验/检查项目列表（如[肝功,肾功,血常规,...]）+ 实际完成项目列表，逐一匹配',
                '现有 arrayContains 仅做单值包含；建议升级为 arrayContainsAll（列表子集检查），或画布上多个 contains 条件 AND 组合')

    # —— 10.12 死亡病例讨论时间 ——
    if '死亡后' in s and ('一周' in s or '7天' in s):
        return ('timeCheck',
                '需要：死亡时间 + 死亡病例讨论记录创建时间 + 阈值(7天)',
                '')

    # —— 10.13 死亡讨论记录内容包括 X ——
    if re.search(r'死亡病例讨论.*(内容|结论)(需|应当?)?(包含|包括)', s):
        return ('contains',
                '需要：死亡病例讨论记录字段 + 关键词（如讨论时间、主持人、死亡诊断、死亡原因）',
                '若是结构化子项检查，建议升级为 sectionCheck')

    # —— 10.13.1 由 X 组织 Y 进行讨论（参与人员校验） ——
    if re.search(r'由.*主任.*组织.*讨论|全部医生.*讨论', s):
        return ('contains',
                '需要：死亡病例讨论记录的参与人员字段 + 科主任姓名 + 在岗医生名册',
                '建议组合：contains 检查记录中是否含科主任姓名 + arrayContainsAll 检查参与人员是否覆盖在岗医生（需扩展 arrayContainsAll）')

    # —— 10.14 每项医嘱只包含一个内容 ——
    if '应当只包含一个内容' in s or '只包含一个' in s:
        return ('expression',
                '需要：医嘱内容文本 + 语义分割规则（多内容检测逻辑）',
                '现有计算符无法处理"一条医嘱不能包含多个诊疗项目"的语义判断；建议升级 expression 计算符或由 HIS 系统侧控制')

    # —— 10.15 手术医嘱时间早于术前讨论 ——
    if re.search(r'医嘱.*开立时间.*早于|签署时间.*早于.*术前讨论', s):
        return ('timeCheck',
                '需要：手术医嘱开立时间/知情同意书签署时间 + 术前讨论时间 + 比较方向(before)',
                '现有 timeCheck 仅支持 timeDiff 在 [min,max] 区间；此规则需 t1 < t2 语义，建议扩展 timeCheck 支持负值 min 或增加 before/after 语义')

    # —— 11. 数值范围（如：血压、体温、年龄区间） ——
    if re.search(r'\d+\s*(～|-|至|到)\s*\d+', s):
        return ('between', '需要：被检查数值字段、min、max', '')

    # —— 默认 ——
    return ('待人工判定', '需要：根据判断依据语义补充', '该条目未匹配到现有计算符模板，建议人工标注或扩展计算符')


def main():
    wb_in = openpyxl.load_workbook(SRC, data_only=True)
    ws_in = wb_in['住院最新核对版']

    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = '住院规则计算符匹配'

    headers = ['序号', '类别', '书写项目', '判断依据', '计算符', '判断数据需求', '升级建议']
    ws_out.append(headers)
    head_font = Font(bold=True, color='FFFFFF')
    head_fill = PatternFill('solid', fgColor='4472C4')
    for col_idx, _ in enumerate(headers, start=1):
        c = ws_out.cell(row=1, column=col_idx)
        c.font = head_font
        c.fill = head_fill
        c.alignment = Alignment(horizontal='center', vertical='center')

    idx = 0
    op_stat = {}
    cur_cat = ''
    cur_item = ''
    for row in ws_in.iter_rows(min_row=2, max_row=ws_in.max_row, values_only=False):
        cat = row[0].value
        item = row[3].value
        dim = row[9].value
        basis = row[10].value
        if cat: cur_cat = str(cat).strip()
        if item: cur_item = str(item).strip()
        if not (dim and '规则质控' in str(dim) and basis):
            continue
        idx += 1
        op, data_req, suggest = classify(str(basis))
        op_stat[op] = op_stat.get(op, 0) + 1
        ws_out.append([idx, cur_cat, cur_item, str(basis).strip(), op, data_req, suggest])

    # 列宽
    widths = {'A': 6, 'B': 16, 'C': 22, 'D': 60, 'E': 18, 'F': 50, 'G': 50}
    for col, w in widths.items():
        ws_out.column_dimensions[col].width = w
    for r in ws_out.iter_rows(min_row=2, max_row=ws_out.max_row):
        for c in r:
            c.alignment = Alignment(wrap_text=True, vertical='top')

    # 统计页
    ws_stat = wb_out.create_sheet('计算符使用统计')
    ws_stat.append(['计算符', '出现次数', '是否现有'])
    for op, n in sorted(op_stat.items(), key=lambda x: -x[1]):
        is_exist = '是' if op in EXIST_OPS else ('否-需扩展' if op != '待人工判定' else '否-未匹配')
        ws_stat.append([op, n, is_exist])
    for c in ws_stat[1]:
        c.font = head_font
        c.fill = head_fill
    ws_stat.column_dimensions['A'].width = 20
    ws_stat.column_dimensions['B'].width = 12
    ws_stat.column_dimensions['C'].width = 16

    wb_out.save(DST)
    print(f'✓ 写入 {DST}, 共 {idx} 条')
    print('— 计算符使用统计 —')
    for op, n in sorted(op_stat.items(), key=lambda x: -x[1]):
        flag = '已有' if op in EXIST_OPS else '需扩展'
        print(f'  {op:<22} {n:>4}  [{flag}]')

if __name__ == '__main__':
    main()
